#!/usr/bin/env python3
"""
Ghép các tab phụ (Cover, Records, Test Report, Legend...) của bộ TC gốc của khách vào file
.xlsx do md_to_xlsx.js sinh ra, để file bàn giao giữ đúng cấu trúc tab như bản gốc.

Quy tắc:
  - Thứ tự tab cuối cùng = thứ tự tab của file khách.
  - Tab trùng tên với sheet trong file đích (sheet TC đã chuẩn hóa) → giữ sheet của file đích.
  - Tab còn lại → chép từ file khách (giá trị, style, ô gộp, độ rộng cột, chiều cao dòng,
    comment, data validation).
  - Sheet của file đích không có trong file khách → đặt cuối, kèm cảnh báo.
  - Công thức ở tab chép sang mà trỏ tới sheet TC (đã thay bằng bản chuẩn hóa, không còn khối
    thống kê cũ) → phải ghi đè bằng --set, nếu không script dừng với lỗi.

Cách dùng:
  python3 scripts/convert_excel/merge_customer_tabs.py <file_dich.xlsx> <file_khach.xlsx> \
      [--set "Sheet!A1=giá trị"]... [--merge "Sheet!C14:E14"]... [--height "Sheet!14=57"]...

  Giá trị --set: YYYY-MM-DD → ngày; số → số; còn lại → text (bắt đầu bằng = là công thức).
  Ô được --set mà chưa có style sẽ lấy style của ô ngay phía trên (tiện khi thêm dòng mới).
"""

import argparse
import datetime as dt
import re
import sys
from copy import copy

import openpyxl
from openpyxl.comments import Comment
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.datavalidation import DataValidation


def parse_value(raw):
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", raw):
        return dt.datetime.strptime(raw, "%Y-%m-%d")
    if re.fullmatch(r"-?\d+", raw):
        return int(raw)
    if re.fullmatch(r"-?\d+\.\d+", raw):
        return float(raw)
    return raw


def split_ref(ref):
    sheet, cell = ref.rsplit("!", 1)
    return sheet.strip("'"), cell


def copy_sheet(src, dst):
    for row in src.iter_rows():
        for c in row:
            if c.__class__.__name__ == "MergedCell":
                continue
            d = dst.cell(row=c.row, column=c.column, value=c.value)
            if c.has_style:
                d.font = copy(c.font)
                d.fill = copy(c.fill)
                d.border = copy(c.border)
                d.alignment = copy(c.alignment)
                d.protection = copy(c.protection)
                d.number_format = c.number_format
            if c.comment:
                d.comment = Comment(c.comment.text, c.comment.author, width=c.comment.width, height=c.comment.height)
    for rng in src.merged_cells.ranges:
        dst.merge_cells(str(rng))
    for key, dim in src.column_dimensions.items():
        dst.column_dimensions[key].width = dim.width
        dst.column_dimensions[key].hidden = dim.hidden
    for key, dim in src.row_dimensions.items():
        if dim.height is not None:
            dst.row_dimensions[key].height = dim.height
        dst.row_dimensions[key].hidden = dim.hidden
    for dv in src.data_validations.dataValidation:
        new = DataValidation(type=dv.type, formula1=dv.formula1, formula2=dv.formula2, operator=dv.operator,
                             allow_blank=dv.allow_blank, showDropDown=dv.showDropDown)
        new.sqref = dv.sqref
        dst.add_data_validation(new)
    dst.sheet_view.showGridLines = src.sheet_view.showGridLines
    dst.sheet_view.zoomScale = src.sheet_view.zoomScale
    dst.freeze_panes = src.freeze_panes
    if src.sheet_properties.tabColor is not None:
        dst.sheet_properties.tabColor = copy(src.sheet_properties.tabColor)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("target")
    ap.add_argument("customer")
    ap.add_argument("--set", action="append", default=[], dest="sets")
    ap.add_argument("--merge", action="append", default=[])
    ap.add_argument("--height", action="append", default=[], help="Chiều cao dòng: Sheet!số_dòng=chiều_cao")
    args = ap.parse_args()

    tgt = openpyxl.load_workbook(args.target)
    cus = openpyxl.load_workbook(args.customer)

    kept = [n for n in cus.sheetnames if n in tgt.sheetnames]
    copied = [n for n in cus.sheetnames if n not in tgt.sheetnames]
    extra = [n for n in tgt.sheetnames if n not in cus.sheetnames]
    if not kept:
        sys.exit(f"❌ Không sheet nào của file đích trùng tên tab của file khách {cus.sheetnames} – kiểm tra dòng <!-- sheet: ... --> trong .md")
    for n in extra:
        print(f"⚠️  Sheet '{n}' của file đích không có trong file khách – đặt ở cuối")

    for n in copied:
        copy_sheet(cus[n], tgt.create_sheet(n))

    tgt._sheets = [tgt[n] for n in cus.sheetnames] + [tgt[n] for n in extra]

    for spec in args.sets:
        ref, raw = spec.split("=", 1)
        sheet, addr = split_ref(ref)
        ws = tgt[sheet]
        cell = ws[addr]
        if not cell.has_style and cell.row > 1:
            above = ws.cell(row=cell.row - 1, column=cell.column)
            if above.has_style:
                cell._style = copy(above._style)
        cell.value = parse_value(raw)
    for spec in args.merge:
        sheet, rng = split_ref(spec)
        tgt[sheet].merge_cells(rng)
    for spec in args.height:
        ref, h = spec.split("=", 1)
        sheet, row = split_ref(ref)
        tgt[sheet].row_dimensions[int(row)].height = float(h)

    # Công thức còn trỏ tới sheet TC đã thay → số liệu sai, bắt buộc --set
    pat = re.compile(r"(?:'?(%s)'?)!" % "|".join(re.escape(n) for n in kept))
    broken = [f"{n}!{c.coordinate}: {c.value}" for n in copied for row in tgt[n].iter_rows() for c in row
              if isinstance(c.value, str) and c.value.startswith("=") and pat.search(c.value)]
    if broken:
        sys.exit("❌ Còn công thức trỏ tới sheet TC đã thay – ghi đè bằng --set:\n  " + "\n  ".join(broken))

    # openpyxl không tính công thức → bắt Excel tính lại khi mở để Test Report hiện đúng số
    if tgt.calculation is None:
        tgt.calculation = CalcProperties()
    tgt.calculation.fullCalcOnLoad = True
    tgt.save(args.target)
    print(f"✅ {args.target}: {len(copied)} tab chép từ file khách ({', '.join(copied)}), thứ tự tab = {tgt.sheetnames}")


if __name__ == "__main__":
    main()
