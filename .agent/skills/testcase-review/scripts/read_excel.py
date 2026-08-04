#!/usr/bin/env python3
"""
Đọc tệp testcase Excel, trích xuất dữ liệu có cấu trúc
"""

import json
import sys
import argparse
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("Lỗi: Chưa cài đặt openpyxl. Hãy chạy: pip install openpyxl")
    sys.exit(1)


def read_excel_cases(file_path: str) -> dict:
    """
    Đọc tệp testcase Excel

    Tham số:
        file_path: Đường dẫn tệp Excel

    Trả về:
        Từ điển chứa danh sách phân hệ (modules) và các ca kiểm thử (testcases)
    """
    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb.active

    # Lấy dòng tiêu đề (header)
    headers = []
    for cell in ws[1]:
        headers.append(cell.value)

    # Ánh xạ tiêu đề đã được chuẩn hóa
    header_map = {}
    for idx, header in enumerate(headers):
        if header:
            header_lower = str(header).strip().lower()
            if '编号' in header_lower or 'id' in header_lower or 'mã' in header_lower:
                header_map['case_id'] = idx
            elif '模块' in header_lower or '所属' in header_lower or 'phân hệ' in header_lower or 'module' in header_lower:
                header_map['module'] = idx
            elif '标题' in header_lower or '名称' in header_lower or 'tiêu đề' in header_lower or 'tên' in header_lower:
                header_map['title'] = idx
            elif '类型' in header_lower or 'loại' in header_lower:
                header_map['type'] = idx
            elif '优先级' in header_lower or '优先' in header_lower or 'độ ưu tiên' in header_lower or 'ưu tiên' in header_lower:
                header_map['priority'] = idx
            elif '步骤' in header_lower or 'bước' in header_lower:
                header_map['steps'] = idx
            elif '预期' in header_lower or 'kết quả' in header_lower or 'mong muốn' in header_lower:
                header_map['expected'] = idx

    # Trích xuất các dòng dữ liệu
    testcases = []
    modules = set()

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row or not any(row):
            continue

        case = {
            'case_id': row[header_map.get('case_id', 0)] if header_map.get('case_id') is not None else f'TC_{row_idx:03d}',
            'module': str(row[header_map.get('module', 1)]).strip() if header_map.get('module') is not None and row[header_map.get('module', 1)] else 'Chưa phân loại',
            'title': str(row[header_map.get('title', 2)]).strip() if header_map.get('title') is not None and row[header_map.get('title', 2)] else '',
            'type': str(row[header_map.get('type', 3)]).strip() if header_map.get('type') is not None and row[header_map.get('type', 3)] else 'Kiểm thử chức năng',
            'priority': str(row[header_map.get('priority', 4)]).strip() if header_map.get('priority') is not None and row[header_map.get('priority', 4)] else 'P2',
            'steps': str(row[header_map.get('steps', 6)]).strip() if header_map.get('steps') is not None and row[header_map.get('steps', 6)] else '',
            'expected': str(row[header_map.get('expected', 7)]).strip() if header_map.get('expected') is not None and row[header_map.get('expected', 7)] else ''
        }

        if case['module']:
            modules.add(case['module'])

        testcases.append(case)

    return {
        'modules': sorted(list(modules)),
        'testcases': testcases
    }


def main():
    parser = argparse.ArgumentParser(description='Đọc tệp testcase Excel')
    parser.add_argument('file_path', help='Đường dẫn tệp Excel')
    parser.add_argument('-o', '--output', help='Đường dẫn xuất tệp JSON')

    args = parser.parse_args()

    if not Path(args.file_path).exists():
        print(f"Lỗi: Không tìm thấy tệp: {args.file_path}")
        sys.exit(1)

    result = read_excel_cases(args.file_path)

    if args.output:
        with open(args.output, 'w', encoding='utf-8') as f:
            json.dump(result, f, ensure_ascii=False, indent=2)
        print(f"Đã ghi kết quả đầu ra vào: {args.output}")
    else:
        print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
